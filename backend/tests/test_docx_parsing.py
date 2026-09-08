"""Word 正文提取的行为守卫（app/utils/docx_blocks.py）。

起因：真实 PRD 把整节内容套在一个 1x1 的排版外框表格里，而旧实现用
doc.paragraphs + doc.tables + cell.text，三者都不下钻单元格内的嵌套表格，
导致 8 行 4 列的需求表被静默丢掉（占该文档正文约 80%）。

测 docx_blocks 而不测 parser_service：后者顶部 import httpx / pdfplumber 与
app.config，CI 只装 pytest，一 import 就挂。
"""
from io import BytesIO

import pytest
from docx import Document as DocxDocument

from app.utils.docx_blocks import render_blocks


def _render(doc) -> list[str]:
    out: list[str] = []
    render_blocks(doc.element.body, doc, out)
    return out


def _roundtrip(doc):
    """存盘再读回，确保测的是真实 docx 的 XML 结构而非内存对象。"""
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return DocxDocument(buf)


def test_单元格内嵌套的表格不丢失():
    doc = DocxDocument()
    doc.add_paragraph("章节正文")
    outer = doc.add_table(rows=1, cols=1)          # 排版外框
    cell = outer.cell(0, 0)
    cell.text = "21.1 元素表"
    inner = cell.add_table(rows=2, cols=2)
    inner.cell(0, 0).text = "元素"
    inner.cell(0, 1).text = "数据来源"
    inner.cell(1, 0).text = "页签"
    inner.cell(1, 1).text = "条目数"

    lines = _render(_roundtrip(doc))

    assert "元素 | 数据来源" in lines
    assert "页签 | 条目数" in lines


def test_1x1外框只透传内容不产生表格行():
    doc = DocxDocument()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "框内标题"
    cell.add_paragraph("框内正文")

    lines = _render(_roundtrip(doc))

    # 透传成独立的两行，而不是被压成一行。
    assert lines == ["框内标题", "框内正文"]


def test_空单元格保留占位以免整行列错位():
    doc = DocxDocument()
    t = doc.add_table(rows=1, cols=4)
    t.cell(0, 1).text = "固定文案"
    t.cell(0, 2).text = "—"
    t.cell(0, 3).text = "—"

    lines = _render(_roundtrip(doc))

    # 首列为空也要占位，否则「固定文案」会左移到第一列。
    assert lines == [" | 固定文案 | — | —"]


def test_相邻重复单元格不去重():
    """真实表格里相邻列常同为「—」，去重会吃掉列。"""
    doc = DocxDocument()
    t = doc.add_table(rows=1, cols=3)
    for i, v in enumerate(["空态", "—", "—"]):
        t.cell(0, i).text = v

    assert _render(_roundtrip(doc)) == ["空态 | — | —"]


def test_横向合并的值只出现一次其余列补空():
    """row.cells 会把同一个 tc 在每个被跨的列各返回一次，重复值会造出假的重复列。"""
    doc = DocxDocument()
    t = doc.add_table(rows=2, cols=3)
    for i, v in enumerate(["元素", "来源", "操作"]):
        t.cell(0, i).text = v
    t.cell(1, 0).merge(t.cell(1, 1))       # 先合并前两列，再赋值：真实文档里跨列格只有一个值
    t.cell(1, 0).text = "合计"
    t.cell(1, 2).text = "100"

    lines = _render(_roundtrip(doc))

    # 值归第一列、被跨的第二列补空，第三列仍对齐到「操作」。
    assert lines == ["元素 | 来源 | 操作", "合计 |  | 100"]


def test_整行横向合并不重复成多列():
    """整行合并常被当分组标题用，重复值会变成「分组A | 分组A | 分组A」。"""
    doc = DocxDocument()
    t = doc.add_table(rows=1, cols=3)
    t.cell(0, 0).text = "一、基础信息"
    t.cell(0, 0).merge(t.cell(0, 1)).merge(t.cell(0, 2))

    assert _render(_roundtrip(doc))[0].count("一、基础信息") == 1


def test_纵向合并的值在每行重复以便每行自解释():
    """fill-down 是想要的行为：分类列跨多行时，每行都该带上分类名。"""
    doc = DocxDocument()
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "分类A"
    t.cell(0, 1).text = "行一"
    t.cell(1, 1).text = "行二"
    t.cell(0, 0).merge(t.cell(1, 0))       # 首列跨两行

    lines = _render(_roundtrip(doc))

    assert len(lines) == 2
    assert all("分类A" in line for line in lines), lines


def test_段落与表格按文档顺序输出():
    doc = DocxDocument()
    doc.add_paragraph("表前段落")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    doc.add_paragraph("表后段落")

    # 旧实现把所有表格追加到所有段落之后，表会脱离它所属的章节标题。
    assert _render(_roundtrip(doc)) == ["表前段落", "A | B", "表后段落"]


@pytest.mark.parametrize("raw, expected", [
    ("尾部换行\n", "尾部换行"),
    ("  两侧留白  ", "两侧留白"),
])
def test_段落文本去除首尾空白(raw, expected):
    doc = DocxDocument()
    doc.add_paragraph(raw)

    assert _render(_roundtrip(doc)) == [expected]
