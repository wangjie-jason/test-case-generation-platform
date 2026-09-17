"""docx 正文的块级遍历与纯文本渲染（只依赖 python-docx）。

为什么不能直接用 python-docx 的现成接口——三个坑叠在一起，表格会静默消失：
  · doc.paragraphs 只返回 body 顶层段落，跳过表格内的段落；
  · doc.tables 只返回 body 顶层表格，不含单元格里嵌套的表格；
  · cell.text 等于 "\\n".join(cell.paragraphs)，只拼直接子段落，嵌套表格直接被忽略。
真实 PRD 常把整节内容套在一个 1x1 的排版外框表格里，于是外框内那张真正的需求表
（连表头一起）全部丢掉，且不报错。故这里改为按文档顺序自己走 XML 并递归下钻。

合并单元格无需担心列错位：row.cells 在任何合并下都返回完整列数（已实测）。横向与纵向
两种合并的处理差异见 render_table 内注释。
"""
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(parent_elm, doc):
    """按文档顺序产出块级对象：Paragraph 或 Table。"""
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def render_blocks(parent_elm, doc, out: list[str]) -> None:
    """把 parent_elm 下的段落与表格按文档顺序追加成文本行。"""
    for block in iter_blocks(parent_elm, doc):
        if isinstance(block, Paragraph):
            if block.text.strip():
                out.append(block.text.strip())
        else:
            render_table(block, doc, out)


def render_table(table: Table, doc, out: list[str]) -> None:
    # 1x1 表格是排版外框而非数据表，透传内容，避免把整节正文压成一行、并保住框内的顺序。
    if len(table.rows) == 1 and len(table.rows[0].cells) == 1:
        render_blocks(table.rows[0].cells[0]._tc, doc, out)
        return

    for row in table.rows:
        texts, nested = [], []
        spanned = set()
        for cell in row.cells:
            # 横向合并（gridSpan）时 row.cells 会把同一个 w:tc 在每个被跨的列各返回一次。
            # 值归第一列，其余列补空占位——不重复值（否则整行合并的分组标题会变成
            # 「一、基础信息 | 一、基础信息 | 一、基础信息」，凭空造出重复列），也不合并成
            # 一格（那会让本行的列数少于其它行、列全部左移）。
            # 只按 tc 身份判定，故合法重复的文本（相邻列同为「—」）不受影响。
            # 纵向合并（vMerge）不进这个分支：被合并的 tc 在每行各出现一次，
            # 值会在各行重复，正是想要的 fill-down——每行自解释，分类才不会丢。
            if cell._tc in spanned:
                texts.append("")
                continue
            spanned.add(cell._tc)

            lines = []
            for block in iter_blocks(cell._tc, doc):
                if isinstance(block, Paragraph):
                    if block.text.strip():
                        lines.append(block.text.strip())
                else:
                    # 嵌套表格不能塞进当前行，整行渲染完再单独输出。
                    nested.append(block)
            texts.append(" ".join(lines))
        # 空单元格要保留占位：过滤掉会让整行的列左移错位，喂给 LLM 就是错事实。
        if any(texts):
            out.append(" | ".join(texts))
        for t in nested:
            render_table(t, doc, out)
